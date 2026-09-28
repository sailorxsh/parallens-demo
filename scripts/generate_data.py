"""Generate a deterministic, internally reconciled dashboard demo dataset."""
from __future__ import annotations

import json
from datetime import date, timedelta
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "prototype" / "data"
BOOK = ROOT / "revised" / "用户与设备经营看板-修订数据底表.xlsx"
DATA.mkdir(parents=True, exist_ok=True)


def write(name: str, value: object) -> None:
    target = DATA / name
    # These files were reconciled with the revised workbook after the base
    # fixture was first generated. Rebuilding synthetic facts must not erase
    # their reviewed metric wording, windows or page evidence.
    if name in {"metric-catalog.json", "filter-contract.json", "diagnostics.json"} and target.exists():
        return
    target.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


workbook = load_workbook(BOOK, data_only=True)
sheet = workbook["框架版底表"]
catalog = []
for cells in sheet.iter_rows(min_row=4, max_row=57, values_only=True):
    (number, block, layer, name, definition, formula, default, grain, source,
     automation, chart, drill, role, question, scope, note, page, home) = cells
    catalog.append({
        "id": number, "name": name, "block": block, "layer": layer,
        "definition": definition, "formula": formula, "default": default,
        "grain": grain, "source": source, "automated": automation,
        "chart": chart, "drill": drill, "homeRole": home,
        "businessQuestion": question, "scope": scope, "note": note,
        "page": page,
    })
write("metric-catalog.json", catalog)

registrations = [2420, 2490, 2510, 2460, 2600, 2580, 2670, 2640, 2700, 2730, 2690, 2750]
weekly = []
first_end = date(2026, 6, 28)
for i, registrations_count in enumerate(registrations):
    late = i >= 9
    firmware_issue = i >= 8
    week_end = first_end + timedelta(weeks=i)
    bind72 = round(registrations_count * (0.72 + (i % 3) * 0.004))
    first_image = round(registrations_count * (0.65 - (0.01 if firmware_issue else 0)))
    bird_eligible = round(registrations_count * 0.70)
    value = round(bird_eligible * (0.56 if not firmware_issue else 0.48))
    trial_start = round(registrations_count * 0.47)
    segment_paid = 178 if late else 194
    other_paid = 256
    empty_k6 = 3280 if firmware_issue else 1680
    empty_others = 2520
    weekly.append({
        "week": i + 1, "weekEnd": week_end.isoformat(),
        "registrationCohort": {"eligible": registrations_count, "registered": registrations_count,
                               "bound72h": bind72, "firstImage7d": first_image,
                               "valueActivated7dBird": value,
                               "birdEligible": bird_eligible,
                               "trialStarted7d": trial_start},
        "bindingAttempts": {"total": 2000, "successful": 1800},
        "trialMaturity": {"completed": 1000, "converted": segment_paid + other_paid,
                          "plusMonthlyAndroid": {"completed": 400, "paymentAttempts": 200,
                                                 "paymentFailed": 22 if late else 6,
                                                 "converted": segment_paid},
                          "otherSegments": {"converted": other_paid}},
        "renewal": {"due": 500, "successful": 430},
        "deviceEventsBird": {"triggers": 10000, "empty": empty_k6 + empty_others,
                             "k6Firmware28": {"triggers": 4000, "empty": empty_k6},
                             "others": {"triggers": 6000, "empty": empty_others}},
        "activity": {"compositeMAU": 64400 + 320 * i + (0 if i < 8 else 80),
                     "appMAU": 43500 + 300 * i,
                     "activeDevices": 82000 + 500 * i - (2222 if firmware_issue else 0)},
        "paidUsersIncludingPackOnly": 12000 + 50 * i,
    })
weekly[-1]["activity"]["compositeMAU"] = 68000
weekly[-1]["activity"]["appMAU"] = 46800
write("weekly.json", weekly)

statuses = {
    "bird": {"free": 47880, "trialEarly": 1000, "trialNearExpiry": 500,
             "paidCurrent": 7400, "cancelledButEntitled": 600, "paymentFailed": 300},
    "hunting": {"noEntitlement": 18120, "trialEarly": 1000, "trialNearExpiry": 500,
                "paidCurrent": 4000, "cancelledButEntitled": 400,
                "paymentFailed": 300, "expired": 400},
}
snapshot = {
    "asOf": "2026-09-24", "latestCompleteWeek": weekly[-1]["weekEnd"],
    "season": "迁徙季", "unusedThresholdDays": 21,
    "registeredCumulative": 128600, "boundOwnerAccounts": 82400,
    "activatedDevicesCumulative": 176200, "boundDevicesCurrent": 143800,
    "subscribersOpening": 12020, "firstPaidFromTrial": 530,
    "firstPaidDirect": 120, "subscriptionRecovered": 90,
    "subscriptionLost": 360, "effectiveSubscribers": 12400,
    "subscriptionNetAdd": 380, "baseMRR": 52000, "packMRR": 4000,
    "totalMRR": 56000, "compositeMAU": 68000, "appMAU": 46800,
    "paidUsersIncludingPackOnly": 12550, "packOnlyUsers": 150,
    "birdOwners": 57680, "huntingOwners": 24720, "subscriptionStatuses": statuses,
    "effectiveActivatedDevices": 139800, "activeDevices": 85278,
    "multiDeviceOwners": 18128,
    "sharedBoundDevices": 25884,
    "alerts": {"boundNotActivated3d": 128, "highChurnRisk": 168,
               "seasonalUnusedOwners": 4200, "freeRightsExhausted": 312},
    "quality": {"eventLossNumerator": 220, "eventExpected": 10000,
                "mappedRecords": 9650, "mappingRecords": 10000,
                "thresholds": {"eventLossMax": .03, "idMappingMin": .95},
                "status": "正常"},
}
write("snapshot.json", snapshot)

latest = weekly[-1]
rate = lambda n, d: {"kind": "rate", "numerator": n, "denominator": d, "value": round(n / d, 6)}
count = lambda value: {"kind": "count", "value": value}
money = lambda value: {"kind": "usd", "value": value}
group = lambda value: {"kind": "group", "value": value}
cohort_30d = {"week": "W7", "eligible": weekly[6]["registrationCohort"]["registered"],
              "firstImage": 2140}
active_exclusive = {
    "appOnly": 23800, "contentOnly": 19200, "pushOnly": 200,
    "appContentOnly": 15000, "appPushOnly": 4000,
    "contentPushOnly": 1800, "allThree": 4000,
}
rights = {"subscriptionOnly": 9400, "packOnly": 150, "both": 3000}
packs = {"Extra Memory Pack": 1200, "Device Pack": 1000,
         "观鸟 AI Agent Pack": 800, "HD Pack": 600, "狩猎 AI Agent Pack": 450}
app_platform = {"iOS": 57870, "Android": 70730}
device_models = {"K6": 16000, "Bird Lite": 50000,
                 "Bird Pro": 45000, "Hunt Pro": 32800}
model_performance = {
    "K6": {"owners": 9000, "activeRate": .62, "subscriptionRate": .17},
    "Bird Lite": {"owners": 24000, "activeRate": .68, "subscriptionRate": .14},
    "Bird Pro": {"owners": 18000, "activeRate": .75, "subscriptionRate": .20},
    "Hunt Pro": {"owners": 15000, "activeRate": .60, "subscriptionRate": .13},
}
channel_performance = {
    "Amazon": {"declarationOnly": True, "owners": 12000,
               "activation7dRate": .64, "activeRate": .70,
               "subscriptionRate": .14, "monthlyPayers": 1000,
               "subscriptionNetIncome": 5100, "ARPPU": 5.10},
    "Shopify": {"declarationOnly": True, "owners": 8000,
                "activation7dRate": .70, "activeRate": .75,
                "subscriptionRate": .19, "monthlyPayers": 800,
                "subscriptionNetIncome": 4320, "ARPPU": 5.40},
}
firmware_counts = {"latest": 100660, "k6Firmware28": 15000,
                   "otherOlder": 28140}
device_count_buckets = [
    {"label": "1台", "owners": 64272, "devices": 64272},
    {"label": "2台", "owners": 4000, "devices": 8000},
    {"label": "3台", "owners": 4000, "devices": 12000},
    {"label": "4台及以上", "owners": 10128, "devices": 59528},
]
unbind = {"events": 300, "observedDevices": 250,
          "reboundWithin7d": 120, "reboundWithin30d": 180}
metrics = {
    2: count(12400), 3: money(56000), 4: count(68000), 5: group(statuses),
    6: rate(12400,82400), 7: group({"registered":128600,"boundOwners":82400}),
    8: group({"activatedCumulative":176200,"boundCurrent":143800}),
    9: count(2750), 10: rate(latest["registrationCohort"]["bound72h"],2750),
    11: rate(1800,2000), 12: group({
        "D7": rate(latest["registrationCohort"]["firstImage7d"],2750),
        "D30": rate(cohort_30d["firstImage"],cohort_30d["eligible"]),
        "D30Cohort": cohort_30d["week"]}),
    13: rate(latest["registrationCohort"]["valueActivated7dBird"],latest["registrationCohort"]["birdEligible"]),
    14: group({"p50Minutes":170,"p90Minutes":1440}), 15: count(128),
    16: count(2450), 17: rate(latest["registrationCohort"]["trialStarted7d"],2750),
    18: rate(434,1000), 19: rate(650,2500), 20: group({
        "combined":rate(120,1400), "subscriptionPage":rate(90,900),
        "popup":rate(30,500)}),
    21: group({"opening":12020,"trialPaid":530,"directPaid":120,"recovered":90,"lost":360,"closing":12400}),
    22: group({"anyExhausted":rate(312,47880), "aiExhausted":rate(220,47880),
               "storageExhausted":rate(160,47880), "transferExhausted":rate(140,47880)}),
    23: rate(430,500), 24: rate(27000,35000),
    25: group({"DAU":12600,"WAU":29800,"MAU":68000,"appMAU":46800}),
    26: group({"sessionsPerUser":2.1,"durationP50Minutes":4.5,"durationP90Minutes":26}),
    27: rate(26000,68000), 28: group({"D1":rate(1900,2500),"D7":rate(1500,2500),"D30":rate(1100,2500)}),
    29: rate(4200,82400),
    30: group({"composite":68000,"app":46800,"content":40000,
               "pushClicked":10000,"exclusiveCells":active_exclusive,"overlapAllowed":True}),
    31: group({"delivery":rate(90000,100000),"click":rate(13500,90000)}),
    32: group({"upgrade":rate(90,8000), "downgrade":rate(60,8000)}),
    33: group({"recovery":rate(90,900), "durationP50Months":8,
               "durationP90Months":24}), 34: count(12550),
    35: group({"monthlySubscriptionPayers":10000,"subscriptionNetIncome":52000,"ARPPU":5.2}),
    36: group({"anyPack":rate(3150,12550),
               "byPack":{name:rate(holders,12550) for name,holders in packs.items()}}),
    37: group({"firstPurchase":400,"repeatOrRenewal":2750}),
    38: group({"countryShare":{"US":0.65,"UK":0.18,"DE":0.10,"Other":0.07},
               "appPlatformShare":{name:count/128600 for name,count in app_platform.items()},
               "deviceModelShare":{name:count/143800 for name,count in device_models.items()}}),
    39: group(model_performance),
    40: group(channel_performance),
    41: rate(85278,139800), 42: rate(5800,10000),
    43: group({"D7":rate(1600,2000),"D30":rate(1450,2000)}),
    44: rate(12000,139800),
    45: group({"networkSuccess":rate(9500,10000),"reportedWithin24h":120000}),
    46: group({"upload":rate(9900,10000),"live":rate(9700,10000),"push":rate(9850,10000)}),
    47: rate(100660,143800), 48: rate(18128,82400), 49: rate(25884,143800),
    50: group({"boundActivated":139800,"boundNotActivated":4000,"activatedUnbound":36400,
               "sharedBoundActivated":25000,"sharedBoundNotActivated":884}),
    51: group({"cohort":"2026-07", "unboundEvents":unbind["events"],
               "observedUnboundDevices":unbind["observedDevices"],
               "unboundDeviceRate":rate(unbind["observedDevices"],142000),
               "reboundWithin7d":rate(unbind["reboundWithin7d"],unbind["observedDevices"]),
               "reboundWithin30d":rate(unbind["reboundWithin30d"],unbind["observedDevices"])}),
    52: count(168), 53: count(7),
    54: group({"eventLoss":rate(220,10000),"idMapping":rate(9650,10000),"status":"正常"}),
}
values = {f"m{index:02}": metrics[index] for index in range(2,55)}
write("metric-values.json", values)
pages = {page: [] for page in "ABCDEFG"}
for item in catalog[1:]:
    pages[item["page"]].append({"metricId":item["id"], "value":values[f'm{item["id"]:02}']})
write("page-data.json", pages)

stories = {
    "A": {"question":"试用转正率为何下降？", "periods":{"baselineWeeks":"W1–W9","issueWeeks":"W10–W12"},
          "baseline":{"trialCompleted":1000,"converted":450,"segmentAttempts":200,"segmentFailed":6,"segmentConverted":194,"othersConverted":256},
          "issue":{"trialCompleted":1000,"converted":434,"segmentAttempts":200,"segmentFailed":22,"segmentConverted":178,"othersConverted":256},
          "segment":"Plus · 月付 · Android", "interpretation":"支付失败增加16人，解释总转正率下降1.6个百分点。"},
    "B": {"question":"价值激活下滑时，设备内容质量有何异常？",
          "periods":{"baselineWeeks":"W1–W8","issueWeeks":"W9–W12"},
          "baseline":{"allTriggers":10000,"allEmpty":4200,"k6Firmware28Triggers":4000,"k6Firmware28Empty":1680,"otherTriggers":6000,"otherEmpty":2520},
          "issue":{"allTriggers":10000,"allEmpty":5800,"k6Firmware28Triggers":4000,"k6Firmware28Empty":3280,"otherTriggers":6000,"otherEmpty":2520},
          "segment":"K6 · 固件 2.8", "interpretation":"设备侧出现集中异常；与用户价值激活下滑同期，仍需排查因果。"},
}
write("stories.json", stories)


def ratio(numerator: int, denominator: int) -> str:
    return f"{numerator / denominator * 100:.1f}%"


def ratio_cell(numerator: int, denominator: int) -> str:
    return f"{numerator:,} / {denominator:,} · {ratio(numerator,denominator)}"


def series(title, unit, labels, lines, note):
    return {"title": title, "unit": unit, "labels": labels,
            "series": [{"name": name, "values": values} for name, values in lines],
            "note": note}


def table(title, columns, rows, note=""):
    return {"title": title, "columns": columns, "rows": rows, "note": note}


def action(code, trigger, suggestion, owner):
    object_type = "演示设备" if code[0] in "ACF" else "演示异常" if code[0] == "G" else "演示主账号"
    return {"id": code, "object": f"{object_type} {code}", "trigger": trigger,
            "suggestion": suggestion, "owner": owner, "status": "待研判",
            "updatedAt": "2026-09-24 09:00（模拟）"}


labels = [f"W{w['week']}" for w in weekly]
cohort = latest["registrationCohort"]
base_a, issue_a = stories["A"]["baseline"], stories["A"]["issue"]
base_b, issue_b = stories["B"]["baseline"], stories["B"]["issue"]
diagnostics = {
    "A": {
        "period": "最新完整周 W12 · 2026-09-13", "heroIds": [9, 10, 11, 12, 13, 15],
        "finding": "观鸟线7日价值激活率从前8周的56.0%降至近4周的48.0%。K6 固件2.8空触发同期上升，但现有数据只支持并发排查。",
        "crossLink": "C", "crossLabel": "继续排查设备内容质量",
        "chart": series("新客批次转化率 · 最近12个完整周", "%", labels, [
            ("72小时绑定率", [round(w["registrationCohort"]["bound72h"] / w["registrationCohort"]["registered"] * 100, 1) for w in weekly]),
            ("7日首图率", [round(w["registrationCohort"]["firstImage7d"] / w["registrationCohort"]["registered"] * 100, 1) for w in weekly]),
            ("观鸟7日价值激活率", [round(w["registrationCohort"]["valueActivated7dBird"] / w["registrationCohort"]["birdEligible"] * 100, 1) for w in weekly]),
        ], "三条率的分母不同；仅比较各自趋势，不连成同一漏斗。"),
        "tables": [table("最新注册批次核对", ["节点", "分子", "分母", "率 / 数量"], [
            ["72小时首绑", cohort["bound72h"], cohort["registered"], ratio(cohort["bound72h"], cohort["registered"])],
            ["7日首图", cohort["firstImage7d"], cohort["registered"], ratio(cohort["firstImage7d"], cohort["registered"])],
            ["观鸟7日价值激活", cohort["valueActivated7dBird"], cohort["birdEligible"], ratio(cohort["valueActivated7dBird"], cohort["birdEligible"])],
            ["7日开启试用", cohort["trialStarted7d"], cohort["registered"], ratio(cohort["trialStarted7d"], cohort["registered"])],
            ["狩猎线7日价值激活", "不适用", "—", "定义待确认"],
        ], "观鸟价值激活使用观鸟合格批次；不把注册、绑定、首图、试用的不同人群强制串接。"),
        table("需跟进的激活阻塞", ["信号", "当前值", "解读"], [
            ["已绑定超过3天仍无首图", snapshot["alerts"]["boundNotActivated3d"], "清单池，需核对真实设备状态"],
            ["绑定→首图耗时 P50 / P90", "170 / 1,440 分钟", "长尾优先排查"],
            ["本周新增激活设备", 2450, "首次产出内容设备"],
        ])],
        "actions": [action("A-01", "绑定超过3天无首图 · 演示样本", "核实联网、固件与首图链路", "设备运营"),
                    action("A-02", "观鸟7日价值激活下降 · 演示样本", "按型号与固件复核空触发", "产品分析")],
    },
    "B": {
        "period": "快照 2026-09-24；趋势为最近12周", "heroIds": [4, 24, 25, 27, 29, 30],
        "finding": "复合MAU为68,000，App MAU为46,800。主动使用、内容行为和推送点击可能重叠，不能相加当作独立用户。",
        "crossLink": "C", "crossLabel": "查看设备行为是否同步变化",
        "chart": series("活跃规模 · 最近12个完整周", "人", labels, [
            ("复合MAU", [w["activity"]["compositeMAU"] for w in weekly]),
            ("App MAU", [w["activity"]["appMAU"] for w in weekly]),
        ], "两条线是滚动30天去重主账号数，App MAU是复合活跃的子集。"),
        "tables": [table("复合活跃来源", ["来源", "去重主账号", "占复合MAU", "关系"], [
            ["复合活跃", 68000, "100.0%", "三类来源并集"],
            ["App主动活跃", 46800, ratio(46800, 68000), "可与其他来源重叠"],
            ["内容活跃", 40000, ratio(40000, 68000), "可与其他来源重叠"],
            ["推送点击活跃", 10000, ratio(10000, 68000), "可与其他来源重叠"],
        ], "来源人数之和不应等于复合MAU；本样本未提供交叉集合明细。"),
        table("留存与回流观察", ["指标", "分子", "分母", "结果"], [
            ["月度留存", 27000, 35000, ratio(27000, 35000)],
            ["新客D1留存", 1900, 2500, ratio(1900, 2500)],
            ["新客D7留存", 1500, 2500, ratio(1500, 2500)],
            ["新客D30留存", 1100, 2500, ratio(1100, 2500)],
            ["迁徙季21天未使用", 4200, 82400, ratio(4200, 82400)],
        ], "不同留存指标使用不同批次；季节阈值是演示设置。")],
        "actions": [action("B-01", "迁徙季21天未使用 · 演示样本", "先排查设备可用性，再设计回流触达", "用户运营"),
                    action("B-02", "复合活跃依赖非App行为 · 演示样本", "核对内容活跃和点击交叉集合", "数据分析")],
    },
    "C": {
        "period": "最新完整周 W12 · 2026-09-13", "heroIds": [8, 41, 42, 43, 45, 46],
        "finding": "观鸟线空触发率由42.0%升至58.0%。K6 固件2.8分组升至82.0%，其余设备保持42.0%；这解释了总体变化，但尚未证明对价值激活的因果影响。",
        "crossLink": "A", "crossLabel": "对照新客价值激活",
        "chart": series("观鸟空触发率 · 最近12个完整周", "%", labels, [
            ("整体", [round(w["deviceEventsBird"]["empty"] / w["deviceEventsBird"]["triggers"] * 100, 1) for w in weekly]),
            ("K6 · 固件2.8", [round(w["deviceEventsBird"]["k6Firmware28"]["empty"] / w["deviceEventsBird"]["k6Firmware28"]["triggers"] * 100, 1) for w in weekly]),
            ("其他设备", [round(w["deviceEventsBird"]["others"]["empty"] / w["deviceEventsBird"]["others"]["triggers"] * 100, 1) for w in weekly]),
        ], "每周触发量：K6 固件2.8为4,000次，其他设备为6,000次；总体为加权汇总。"),
        "tables": [table("空触发异常贡献", ["周期 / 分组", "空触发", "触发", "空触发率"], [
            ["W1–W8 · K6 固件2.8", 1680, 4000, "42.0%"],
            ["W9–W12 · K6 固件2.8", 3280, 4000, "82.0%"],
            ["W1–W8 · 其他设备", 2520, 6000, "42.0%"],
            ["W9–W12 · 其他设备", 2520, 6000, "42.0%"],
            ["W1–W8 · 整体", 4200, 10000, "42.0%"],
            ["W9–W12 · 整体", 5800, 10000, "58.0%"],
        ], "总体上升16个百分点，来自K6固件2.8每周新增1,600次空触发 / 10,000次总体触发。"),
        table("设备健康补充信号", ["指标", "分子", "分母", "结果"], [
            ["活跃设备", 85278, 139800, ratio(85278,139800)],
            ["新设备D7留存", 1600, 2000, "80.0%"],
            ["新设备D30留存", 1450, 2000, "72.5%"],
            ["联网成功", 9500, 10000, "95.0%"],
        ], "设备活跃排除仅有心跳的设备；D7/D30是独立新设备批次。")],
        "actions": [action("C-01", "K6 固件2.8空触发集中偏高 · 演示样本", "核验事件分类与固件日志", "设备产品"),
                    action("C-02", "长时间无有效设备行为 · 演示样本", "排查联网、供电与上传链路", "设备运营")],
    },
    "D": {
        "period": "快照 2026-09-24；试用趋势至 W12", "heroIds": [2, 3, 18, 21, 23, 52],
        "finding": "试用转正率从前9周45.0%降至近3周43.4%。Plus · 月付 · Android 的支付失败每周多16人，其他分组转正人数不变，恰好解释1.6个百分点下降。",
        "crossLink": "E", "crossLabel": "查看付费价值结构",
        "chart": series("试用成熟批次 · 最近12个完整周", "%", labels, [
            ("整体试用转正率", [w["trialMaturity"]["converted"] / w["trialMaturity"]["completed"] * 100 for w in weekly]),
            ("Plus月付Android支付失败率", [w["trialMaturity"]["plusMonthlyAndroid"]["paymentFailed"] / w["trialMaturity"]["plusMonthlyAndroid"]["paymentAttempts"] * 100 for w in weekly]),
        ], "两条率使用不同分母；后者为该分组支付尝试次数，不能视为总体转正率的子集比例。"),
        "tables": [table("有效订阅期初到期末", ["流转", "人数", "影响"], [
            ["期初有效订阅", 12020, "起点"],
            ["试用后首次付费", 530, "+"],
            ["直接首次付费", 120, "+"],
            ["恢复订阅", 90, "+"],
            ["退出有效订阅", 360, "−"],
            ["期末有效订阅", 12400, "净增 +380"],
        ], "观鸟Free和试用均不计入有效订阅；已取消但权益未到期者仍计入。"),
        table("支付失败对试用转正的解释", ["指标", "W1–W9（每周）", "W10–W12（每周）", "变化"], [
            ["到期试用人数", 1000, 1000, "持平"],
            ["整体转正", 450, 434, "−16"],
            ["Plus月付Android支付尝试", 200, 200, "持平"],
            ["该组支付失败", 6, 22, "+16"],
            ["该组转正", 194, 178, "−16"],
            ["其他分组转正", 256, 256, "持平"],
        ], "演示分组完整闭合；失败增加16人 = 总转正减少16人。")],
        "actions": [action("D-01", "Plus月付Android支付失败升高 · 演示样本", "核对支付错误码与扣费链路", "订阅产品"),
                    action("D-02", "流失高危主账号 · 演示样本", "复核风险规则和可触达状态", "用户运营"),
                    action("D-03", "观鸟Free权益用满 · 演示样本", "核对权益状态与升级入口", "商业化运营")],
    },
    "E": {
        "period": "权益快照 2026-09-24；付款趋势为最近3个完整月（至8月）", "heroIds": [3, 34, 35, 36, 37, 39],
        "finding": "当前拥有付费权益的主账号12,550人，其中150人仅持有增值Pack。基础订阅MRR为$52,000，Pack MRR为$4,000；月内实际付款人数与期末权益人数不可混用。",
        "crossLink": "D", "crossLabel": "回看订阅流转",
        "chart": series("月内订阅付款主账号 · 最近3个完整月", "人", ["6月", "7月", "8月"], [
            ("付款主账号", [9600, 9800, 10000]),
        ], "按自然月内实际发生订阅付款的主账号去重；与期末付费权益人数不同。"),
        "extraCharts": [series("订阅ARPPU · 相同3个完整月", "美元/人", ["6月", "7月", "8月"], [
            ("ARPPU", [5.00, 5.10, 5.20]),
        ], "仅基础订阅净收入 / 月内订阅付款主账号；与上图共用月份轴，不使用双Y轴。")],
        "tables": [table("收入与权益口径", ["指标", "分子 / 组成", "分母", "结果"], [
            ["期末有效订阅", 12400, "—", "12,400人"],
            ["期末仅持Pack", 150, "—", "150人"],
            ["期末付费权益主账号", "12,400 + 150", "—", "12,550人"],
            ["增值包持有人占比", 3150, 12550, ratio(3150,12550)],
            ["基础订阅MRR", "$52,000", "—", "$52,000"],
            ["增值Pack MRR", "$4,000", "—", "$4,000"],
            ["总MRR", "$52,000 + $4,000", "—", "$56,000"],
            ["月内订阅ARPPU", "$52,000", "10,000人", "$5.20"],
        ], "ARPPU仅用基础订阅净收入 / 月内订阅付款主账号；MRR与当月净收入在本模拟样本数值相同，但概念不同。"),
        table("完整月付款与人均收入", ["月份", "订阅付款主账号", "订阅净收入", "ARPPU"], [
            ["2026-06", 9600, "$48,000", "$5.00"],
            ["2026-07", 9800, "$49,980", "$5.10"],
            ["2026-08", 10000, "$52,000", "$5.20"],
        ], "三个月均为完整自然月；人均收入按本行收入除以本行付款主账号回算。"),
        table("分层观察与证据边界", ["维度", "样本值", "可用判断"], [
            ["增值服务首购 / 复购", "400 / 2,750人", "复购含续购，按主账号去重"],
            ["全部用户国家构成", "美国65% / 英国18% / 德国10% / 其他7%", "不等于付费用户国家构成"],
            ["K6设备型号", "9,000主账号；活跃率62%；订阅率17%", "仅单一型号样本"],
            ["Amazon / Shopify订阅率", "14% / 19%", "仅自报渠道，不支持归因"],
        ], "国家结构与渠道字段均标明自身范围；不据此推出付费贡献或渠道因果。")],
        "actions": [action("E-01", "Pack权益持有差异 · 演示样本", "核对Pack购买与权益生效映射", "商业化运营"),
                    action("E-02", "渠道自报差异 · 演示样本", "先核验来源字段再比较用户质量", "数据分析")],
    },
    "F": {
        "period": "设备与账号快照 2026-09-24", "heroIds": [7, 8, 47, 48, 49, 50],
        "finding": "当前绑定设备143,800台，其中139,800台曾激活，4,000台未激活；历史激活设备176,200台，其中36,400台当前未绑定。两个存量维度不能直接相减解释解绑。",
        "crossLink": "C", "crossLabel": "查看设备健康",
        "chart": series("设备交叉状态 · 当前快照", "台", ["绑定且已激活", "绑定未激活", "已激活未绑定"], [
            ("设备数", [139800, 4000, 36400]),
        ], "三类状态为演示中的交叉状态池；不能与全部出厂设备总量等同。"),
        "tables": [table("设备状态双向核对", ["核对关系", "左侧", "右侧", "闭合值"], [
            ["当前绑定", "已激活 139,800", "未激活 4,000", "143,800台"],
            ["历史激活", "当前绑定 139,800", "当前未绑定 36,400", "176,200台"],
        ], "历史激活与当前绑定是两个维度；任何差值需经交叉状态表解释。纯SD卡且从未接入App的设备不在可观测池。"),
        table("主账号与共享", ["指标", "分子", "分母", "结果"], [
            ["多设备主账号", 18128, 82400, ratio(18128,82400)],
            ["开启共享的绑定设备", 25884, 143800, ratio(25884,143800)],
            ["最新固件覆盖设备", 100660, 143800, ratio(100660,143800)],
            ["解绑事件", 300, "—", "事件数，不是去重设备数"],
            ["7日内重新绑定", 120, 300, ratio(120,300)],
            ["30日内重新绑定", 180, 300, ratio(180,300)],
        ], "7日内重新绑定是30日内重新绑定的子集；解绑事件不能直接解释流失用户。")],
        "actions": [action("F-01", "已绑定未激活 · 演示样本", "核对设备首次产出和绑定时间", "设备运营"),
                    action("F-02", "解绑后未回绑 · 演示样本", "区分转卖、退货与正常换绑", "用户研究")],
    },
    "G": {
        "period": "监控快照 2026-09-24", "heroIds": [53, 54, 15, 29, 42, 52],
        "finding": "当前演示异常事件7条。埋点丢失率2.2%，ID-Mapping覆盖率96.5%，样本标记为“正常”；阈值尚属演示设置，不能据此判断生产数据已达标。",
        "crossLink": "A", "crossLabel": "回看新客异常入口",
        "chart": series("演示异常事件归属", "条", ["新客", "订阅", "设备", "数据质量"], [
            ("事件数", [2, 2, 2, 1]),
        ], "四类事件合计7条；这是模拟事件分配，不是各类告警池人数的和。"),
        "tables": [table("异常事件样本", ["编号", "领域", "信号", "研判状态"], [
            ["EX-01", "新客", "7日价值激活下降", "待核验"],
            ["EX-02", "新客", "绑定后无首图", "待核验"],
            ["EX-03", "订阅", "试用转正下降", "已定位演示贡献"],
            ["EX-04", "订阅", "支付失败上升", "待核验"],
            ["EX-05", "设备", "K6固件2.8空触发升高", "待核验"],
            ["EX-06", "设备", "长期未使用", "待核验"],
            ["EX-07", "数据质量", "ID映射覆盖待观察", "待核验"],
        ], "一行代表一个演示异常事件；同一用户或设备可能出现在多个信号中。"),
        table("数据质量与业务信号", ["指标", "分子", "分母", "结果"], [
            ["埋点事件丢失率", 220, 10000, "2.2%"],
            ["ID-Mapping覆盖率", 9650, 10000, "96.5%"],
            ["事件丢失演示正常阈值", "—", "—", "≤3.0%"],
            ["ID映射演示正常阈值", "—", "—", "≥95.0%"],
            ["数据状态", "演示判断", "—", "正常"],
            ["绑定未激活设备", 128, "—", "清单池"],
            ["季节阈值未使用主账号", 4200, "—", "清单池"],
            ["流失高危主账号", 168, "—", "清单池"],
        ], "不同告警池可能重叠，且事件数与人数/设备数不同，不可相加。")],
        "actions": [action("G-01", "埋点丢失与映射覆盖 · 演示样本", "核对日志抽样与映射规则", "数据工程"),
                    action("G-02", "设备与新客信号同期变化 · 演示样本", "串联批次、设备型号与事件记录", "产品分析")],
    },
}

# Secondary evidence stays separate from the primary chart, but uses the same
# deterministic source quantities as the metric catalog and homepage.
diagnostics["A"]["tables"].extend([
    table("不同完整观察窗的注册批次", ["观察窗", "入组批次", "首图用户", "合格注册", "激活率"], [
        ["7日", "W12", cohort["firstImage7d"], cohort["registered"],
         ratio(cohort["firstImage7d"], cohort["registered"])],
        ["30日", cohort_30d["week"], cohort_30d["firstImage"],
         cohort_30d["eligible"], ratio(cohort_30d["firstImage"], cohort_30d["eligible"])],
    ], "30日率仅取已满30日观察期的 W7 批次；不可把W12的7日分子与W7的30日分母相减。"),
    table("首次绑定失败原因 · W12", ["失败原因", "尝试次数", "占失败尝试", "下一步"], [
        ["蓝牙配对", 85, ratio(85,200), "核对配对错误码"],
        ["WiFi连接", 70, ratio(70,200), "核对网络环境"],
        ["App流程", 45, ratio(45,200), "核对版本与引导"],
        ["合计", 200, "100.0%", "总尝试2,000次，成功1,800次"],
    ], "失败原因按每次首次尝试互斥归类；绑定尝试不是注册批次。"),
])
diagnostics["A"]["extraCharts"] = [{
    "kind":"boxplot", "title":"绑定→首图时长 · 演示分布", "unit":"分钟",
    "labels":["绑定→首图"],
    "series":[{"name":"时长分位", "values":[[20,70,170,480,2400]]}],
    "p90":1440,
    "note":"样本 P50=170分钟、P90=1,440分钟；最小值/Q1/Q3/最大值为额外模拟分布值，不由P50/P90唯一推导。",
}]

diagnostics["B"]["tables"][0]["note"] = (
    "来源人数可重叠；下表将三类活跃拆为7个互斥单元，合计回到68,000人。")
diagnostics["B"]["tables"].extend([
    table("三类活跃来源交叉结构", ["互斥单元", "主账号", "占复合MAU"], [
        [name, active_exclusive[key], ratio(active_exclusive[key],68000)]
        for key,name in [
            ("appOnly","仅App主动"),("contentOnly","仅设备内容"),
            ("pushOnly","仅推送点击"),("appContentOnly","App+内容"),
            ("appPushOnly","App+推送"),("contentPushOnly","内容+推送"),
            ("allThree","三者皆有")]
    ] + [["并集",sum(active_exclusive.values()),"100.0%"]],
    "七个单元互斥；App 46,800、内容 40,000、点击 10,000 均可从这些单元回算。"),
    table("完整月与新客留存批次", ["批次 / 观察期", "留存", "入组", "留存率"], [
        ["2026-06→07 · 月度复合活跃",25200,35000,ratio(25200,35000)],
        ["2026-07→08 · 月度复合活跃",27000,35000,ratio(27000,35000)],
        ["新客样本 · D1",1900,2500,ratio(1900,2500)],
        ["新客样本 · D7",1500,2500,ratio(1500,2500)],
        ["新客样本 · D30",1100,2500,ratio(1100,2500)],
        ["前期未使用池 · 后续回流",600,5000,ratio(600,5000)],
    ], "月度与D1/D7/D30为不同批次；回流率用前期已满观察窗的5,000人，不能拿当前4,200人作分母。"),
    table("推送链路", ["步骤", "分子", "分母", "结果"], [
        ["送达",90000,100000,ratio(90000,100000)],
        ["送达后点击",13500,90000,ratio(13500,90000)],
    ], "推送点击活跃的10,000人为去重主账号；13,500为点击事件，单位不同。"),
    table("使用深度与核心功能", ["指标", "样本值", "解释"], [
        ["复合DAU / WAU / MAU", "12,600 / 29,800 / 68,000人", "各自滚动窗口去重"],
        ["App人均打开次数", "2.1次", "活跃用户使用频次"],
        ["单次使用时长P50 / P90", "4.5 / 26分钟", "时长分位数，非平均值"],
        ["核心功能使用主账号", "26,000 / 68,000人", "38.2%，复合活跃为分母"],
    ], "使用次数与时长的对象为App使用者；功能渗透以复合活跃主账号为分母。"),
])

diagnostics["C"]["tables"].extend([
    table("设备功能与版本", ["诊断项", "成功 / 覆盖", "分母", "结果"], [
        ["24小时有上报",120000,143800,ratio(120000,143800)],
        ["事件上传成功",9900,10000,"99.0%"],
        ["直播连接成功",9700,10000,"97.0%"],
        ["推送功能成功",9850,10000,"98.5%"],
        ["最新固件覆盖",firmware_counts["latest"],143800,
         ratio(firmware_counts["latest"],143800)],
        ["长期未使用设备",12000,139800,ratio(12000,139800)],
        ["前期闲置设备重新活跃",600,3000,ratio(600,3000)],
    ], "上报量以当前绑定设备为观察池；长期未使用率以有效已激活设备为分母；重新活跃另取前期完整观察池。"),
    table("当前可升级设备固件结构", ["版本分组", "设备", "可升级设备", "占比"], [
        ["最新版本",firmware_counts["latest"],143800,
         ratio(firmware_counts["latest"],143800)],
        ["K6 固件2.8",firmware_counts["k6Firmware28"],143800,
         ratio(firmware_counts["k6Firmware28"],143800)],
        ["其他旧版本",firmware_counts["otherOlder"],143800,
         ratio(firmware_counts["otherOlder"],143800)],
    ], "版本设备存量与每周空触发事件量不是同一单位，不能相减。"),
])
diagnostics["C"]["actions"].append(
    action("C-03", "旧固件覆盖 · 演示样本", "先核验适用版本与设备风险，再评估OTA方案", "设备产品"))

diagnostics["D"]["tables"].extend([
    table("转化与续费环节", ["环节", "分子", "分母", "结果"], [
        ["主账号当前有效订阅",12400,82400,ratio(12400,82400)],
        ["新客7日开启试用",cohort["trialStarted7d"],cohort["registered"],
         ratio(cohort["trialStarted7d"],cohort["registered"])],
        ["首次订阅",650,2500,ratio(650,2500)],
        ["订阅页购买",90,900,ratio(90,900)],
        ["运营弹窗购买",30,500,ratio(30,500)],
        ["到期续订",430,500,ratio(430,500)],
        ["期初有效订阅退出",360,12020,ratio(360,12020)],
        ["套餐升级",90,8000,ratio(90,8000)],
        ["套餐降级",60,8000,ratio(60,8000)],
        ["取消或过期后恢复",90,900,ratio(90,900)],
        ["连续订阅时长P50 / P90", "8 / 24个月", "—", "分位数，非平均值"],
    ], "各环节基于自身完整批次；续订未成功70人不等于本期退出有效权益的360人。"),
    table("观鸟Free权益用满", ["权益项", "去重主账号", "Free主账号", "用满率"], [
        ["AI识别30次/月",220,47880,ratio(220,47880)],
        ["存储1GB",160,47880,ratio(160,47880)],
        ["影像传输20次/月",140,47880,ratio(140,47880)],
        ["任一权益用满（并集）",312,47880,ratio(312,47880)],
    ], "三项权益可由同一人同时用满，因此220+160+140不等于并集312。"),
    table("分产品线订阅状态", ["产品线", "Free/无权益", "试用中", "付费中", "取消未到期", "支付失败", "已过期", "合计"], [
        ["观鸟",47880,1500,7400,600,300,0,57680],
        ["狩猎",18120,1500,4000,400,300,400,24720],
    ], "两线内部状态互斥、各自闭合；已取消未到期仍计入有效付费权益。"),
])

rights_chart = series("付费权益互斥结构 · 2026-09-24", "%", ["付费权益主账号"], [
    ("仅订阅",[round(rights["subscriptionOnly"]/12550*100,1)]),
    ("仅Pack",[round(rights["packOnly"]/12550*100,1)]),
    ("两者都有",[round(rights["both"]/12550*100,1)]),
], "人数为9,400 / 150 / 3,000；三类互斥，合计12,550人。")
rights_chart["kind"] = "stacked100"
diagnostics["E"]["extraCharts"].append(rights_chart)
diagnostics["E"]["extraCharts"].append({
    "kind":"heatmap", "title":"设备型号表现矩阵 · 演示样本", "unit":"%",
    "labels":list(model_performance),
    "dimensions":["30天活跃率","主账号订阅率"],
    "series":[{"name":"比率矩阵", "values":[
        [round(values["activeRate"]*100,1),round(values["subscriptionRate"]*100,1)]
        for values in model_performance.values()]}],
    "note":"只比较四个演示型号；两列均为主账号比率，颜色越深比率越高，各格完整分子/分母见下方型号表。",
})
diagnostics["E"]["tables"][2] = table(
    "增值服务首购与复购", ["类型", "去重主账号", "占本期Pack持有人"], [
        ["首次购买",400,ratio(400,3150)],
        ["再次购买或续购",2750,ratio(2750,3150)],
        ["本期Pack持有人",3150,"100.0%"],
    ], "此演示期将Pack持有人互斥归类为首购或再次购买；不与各Pack类型持有人数相加。")
diagnostics["E"]["tables"].extend([
    table("权益与Pack持有结构", ["权益或Pack", "主账号", "分母", "占比"], [
        ["仅订阅",rights["subscriptionOnly"],12550,ratio(rights["subscriptionOnly"],12550)],
        ["仅增值Pack",rights["packOnly"],12550,ratio(rights["packOnly"],12550)],
        ["订阅和Pack都有",rights["both"],12550,ratio(rights["both"],12550)],
        *[[name,holders,12550,ratio(holders,12550)] for name,holders in packs.items()],
    ], "前三项互斥且闭合；各Pack可叠加，五项持有人数不可相加为独立付费用户。"),
    table("设备型号订阅表现 · 主账号可跨型号", ["型号", "主账号", "30天活跃（分子/分母）", "当前订阅（分子/分母）"], [
        [name,values["owners"],
         ratio_cell(round(values["owners"]*values["activeRate"]),values["owners"]),
         ratio_cell(round(values["owners"]*values["subscriptionRate"]),values["owners"])]
        for name,values in model_performance.items()
    ], "同一主账号可拥有多个型号；各型号主账号数不可相加。未提供可支持收入归因的型号映射。"),
    table("自报购买渠道质量 · 样本", ["申报渠道", "主账号", "7日激活（分子/分母）", "30天活跃（分子/分母）", "订阅（分子/分母）", "月内ARPPU（收入/付款人）"], [
        [name,values["owners"],
         ratio_cell(round(values["owners"]*values["activation7dRate"]),values["owners"]),
         ratio_cell(round(values["owners"]*values["activeRate"]),values["owners"]),
         ratio_cell(round(values["owners"]*values["subscriptionRate"]),values["owners"]),
         f'${values["subscriptionNetIncome"]:,} / {values["monthlyPayers"]:,}人 · ${values["ARPPU"]:.2f}']
        for name,values in channel_performance.items()
    ], "渠道仅为用户自报；示例分组未控制国家、型号和批次，不能用作投放因果归因。"),
])
diagnostics["E"]["structureViews"] = [
    {"label":"国家 / 全部注册用户", "unit":"人", "denominator":128600,
     "rows":[[name,round(128600*share)] for name,share in
             [("美国",.65),("英国",.18),("德国",.10),("其他",.07)]],
     "note":"全部注册用户的国家构成，不等于付费用户国家构成。"},
    {"label":"App平台 / 全部注册用户", "unit":"人", "denominator":128600,
     "rows":[[name,number] for name,number in app_platform.items()],
     "note":"平台注册用户按 iOS / Android 互斥归类。"},
    {"label":"设备型号 / 当前绑定设备", "unit":"台", "denominator":143800,
     "rows":[[name,number] for name,number in device_models.items()],
     "note":"设备型号按设备数占比，与主账号数分母不同。"},
]

diagnostics["F"]["extraCharts"] = [series(
    "主账号当前绑定设备数分布", "人", [bucket["label"] for bucket in device_count_buckets],
    [("主账号",[bucket["owners"] for bucket in device_count_buckets])],
    "互斥主账号分层合计82,400人，设备量加权合计143,800台。")]
diagnostics["F"]["tables"][1] = table("主账号、共享与回绑", ["指标", "分子", "分母", "结果"], [
    ["平均绑定设备数",143800,82400,"1.75台/主账号"],
    ["多设备主账号",18128,82400,ratio(18128,82400)],
    ["开启共享的绑定设备",25884,143800,ratio(25884,143800)],
    ["通过共享获得访问权限的用户",30200,"—","去重共享用户"],
    ["最新固件覆盖设备",100660,143800,ratio(100660,143800)],
    ["2026-07解绑事件",unbind["events"],"—","事件数，非去重设备"],
    ["2026-07可观察解绑设备",unbind["observedDevices"],"—","去重设备"],
    ["2026-07解绑设备率",unbind["observedDevices"],142000,
     ratio(unbind["observedDevices"],142000)],
    ["7日内重新绑定",unbind["reboundWithin7d"],unbind["observedDevices"],
     ratio(unbind["reboundWithin7d"],unbind["observedDevices"])],
    ["30日内重新绑定",unbind["reboundWithin30d"],unbind["observedDevices"],
     ratio(unbind["reboundWithin30d"],unbind["observedDevices"])],
], "2026-07解绑批次已满30日观察窗；解绑设备率分母为该月期初绑定设备142,000台。回绑率以同批250台去重设备为分母，7日回绑是30日回绑的子集。")
diagnostics["F"]["tables"].extend([
    table("每主账号绑定设备数分层", ["设备数", "主账号", "设备", "主账号占比"], [
        [bucket["label"],bucket["owners"],bucket["devices"],
         ratio(bucket["owners"],82400)] for bucket in device_count_buckets
    ], "主账号互斥分层；4台及以上为聚合组，不表示每人恰好4台。"),
    table("共享作为绑定状态子维度", ["状态", "有共享", "无共享", "合计"], [
        ["绑定且已激活",25000,114800,139800],
        ["绑定未激活",884,3116,4000],
        ["绑定设备合计",25884,117916,143800],
    ], "共享不另列为第四种绑定状态，避免与已激活/未激活重复计算。"),
])

diagnostics["G"]["tables"].append(table(
    "异常信号与前期基线", ["信号", "前期 / 对照", "最新", "判断边界"], [
        ["观鸟7日价值激活率", "W1–W8 56.0%", "W9–W12 48.0%", "并发线索"],
        ["试用转正率", "W1–W9 45.0%", "W10–W12 43.4%", "演示支付失败贡献闭合"],
        ["观鸟空触发率", "W1–W8 42.0%", "W9–W12 58.0%", "K6固件2.8组解释总量变化"],
        ["埋点丢失率", "7日前 1.8%", "2.2%", "演示阈值，不推断生产达标"],
    ], "业务异常与数据质量监控分开研判；近4周对照示意不等于生产告警规则。"))
diagnostics["G"]["extraCharts"] = [series(
    "埋点事件丢失率 · 最近7日", "%", ["9/18","9/19","9/20","9/21","9/22","9/23","9/24"],
    [("事件丢失率",[1.8,1.9,2.1,2.0,2.2,2.1,2.2])],
    "有可靠预期事件数的演示样本；无预期数时应显示待核验，而非绿色正常。")]
diagnostics["C"]["chart"]["series"][1]["color"] = "#e5484d"
diagnostics["D"]["extraCharts"] = [
    series("Plus月付Android支付失败率 · 最近12个完整周", "%",
           diagnostics["D"]["chart"]["labels"],
           [("该组支付失败率",diagnostics["D"]["chart"]["series"][1]["values"])],
           "分母为该组每周200次支付尝试；与整体试用到期人数1,000人不同。")
]
diagnostics["D"]["extraCharts"][0]["series"][0]["color"] = "#e5484d"
diagnostics["D"]["chart"]["series"] = diagnostics["D"]["chart"]["series"][:1]
diagnostics["D"]["chart"]["note"] = "分母为每周1,000名到期试用用户；下图另看支付尝试失败率，不共用分母。"
diagnostics["G"]["extraCharts"][0]["series"][0]["color"] = "#e5484d"
write("diagnostics.json", diagnostics)

# Explicit coverage for the synthetic slice model. These are supported demo joins,
# not a claim that the source workbook contains row-level observations.
dimension_ids = {
    "country": set(range(2, 54)) - {11, 14, 20, 26, 31, 46, 47, 54},
    "appPlatform": {2,3,4,5,6,7,9,10,11,12,13,17,18,19,20,21,23,24,25,26,27,28,29,30,31,32,33,34,35,36,38,40,48,49,50,51,52,53},
    "productLine": set(range(2, 54)) - {7,9,10,11,14,24,26,28,31,38,48,51,54},
    "deviceModel": {4,6,8,11,13,14,15,16,19,25,27,29,39,40,41,42,43,44,45,46,47,48,49,50,52,53},
    "deviceStatus": {8,11,14,15,16,27,39,41,42,43,44,45,46,47,48,49,50,51,53},
    "season": {13,29,42,44,53},
    "salesChannel": {16,40,43},
    "subscriptionPlatform": {2,3,5,6,17,18,19,21,23,32,33,34,35,36,52},
    "plan": {2,3,5,6,17,18,19,21,22,23,32,33,34,35,36,52},
    "billingCycle": {18},
    "firmware": {41,42,45,46,47,53},
    "functionType": {27,46},
    "anomalyType": {53},
    "week": {9,10,11,12,13,17,18,23,42},
}
windows = {
    "week": "最近完整周及最近12周趋势",
    "day": "2026-09-24 快照或指标定义中的日窗口",
    "month": "最近完整月或指标定义中的月窗口",
    "cohort": "完整观察窗内的指定批次",
}
week_ids = dimension_ids["week"]
month_ids = {3,6,21,22,24,28,35,36,37,38,39,40,43,48,51}
cohort_ids = {10,12,13,17,18,23,28,43,51}
contract = []
for item in catalog[1:]:
    mid = item["id"]
    measure = values[f"m{mid:02d}"]
    entity = str(item["grain"]).split("×")[0].strip()
    if mid in cohort_ids:
        window = windows["cohort"]
    elif mid in week_ids:
        window = windows["week"]
    elif mid in month_ids:
        window = windows["month"]
    else:
        window = windows["day"]
    if measure["kind"] == "rate":
        numerator, denominator = measure["numerator"], measure["denominator"]
    elif measure["kind"] == "group":
        numerator = denominator = "各子项独立，见原指标计算方法"
    else:
        numerator, denominator = "去重计数或金额", "不适用"
    dimensions = [name for name, ids in dimension_ids.items() if mid in ids]
    if mid in {14,26,33,38}:
        dimensions = []
    elif mid == 39:
        dimensions = ["productLine", "deviceModel"]
    elif mid == 40:
        dimensions = ["salesChannel"]
    elif mid == 21:
        dimensions = [name for name in dimensions if name != "productLine"]
    elif mid == 5:
        dimensions = [name for name in dimensions if name != "plan"]
    unsupported = "缺少明确归属关系的组合、未成熟批次、无样本分母或维度未列入本项适用表时显示不适用；不以0代替。"
    if mid in {13,22,42}:
        unsupported += "仅观鸟线有当前口径，狩猎线不适用。"
    if mid == 21:
        unsupported += "缺少分产品线的期初/进入/退出流转。"
    if mid == 5:
        unsupported += "缺少按套餐拆分的互斥状态人数。"
    if mid == 22:
        unsupported += "仅观鸟Free权益，Plus/Pro套餐不适用。"
    contract.append({
        "id": mid, "name": item["name"], "page": item["page"],
        "entity": entity, "window": window,
        "dimensions": dimensions,
        "numerator": numerator, "denominator": denominator,
        "formula": item["formula"], "source": item["source"],
        "rateComponents": [],
        "unsupported": unsupported,
        "model": "模拟独立维度分片；交叉筛选只供原型交互验收，不代表实测分布。",
    })
    def collect_rates(value, path=""):
        if isinstance(value, dict):
            if value.get("kind") == "rate":
                contract[-1]["rateComponents"].append({
                    "name": path or item["name"],
                    "numerator": value["numerator"],
                    "denominator": value["denominator"],
                })
            else:
                for key, child in value.items():
                    collect_rates(child, f"{path}/{key}" if path else key)
    collect_rates(measure)
write("filter-contract.json", contract)
slice_weights = {
    "country": {"US": .65, "UK": .18, "DE": .10, "Other": .07},
    "appPlatform": {"iOS": .45, "Android": .55},
    "productLine": {"Bird": .70, "Hunting": .30},
    "deviceModel": {"K6": .1113, "Bird Lite": .3477, "Bird Pro": .3129, "Hunt Pro": .2281},
    "deviceStatus": {"Effective": .90, "Inactive": .10},
    "season": {"Migration": .70, "Other": .30},
    "salesChannel": {"Amazon": .60, "Shopify": .40},
    "subscriptionPlatform": {"App Store": .44, "Google Play": .36, "Web": .20},
    "plan": {"Free": .54, "Starter": .04, "Plus": .30, "Pro": .12},
    "billingCycle": {"monthly": .8, "annual": .2},
    "firmware": {"2.8": .24, "Other": .76},
    "functionType": {"Upload": .46, "Recognition": .34, "Push": .20},
    "anomalyType": {"Business": .57, "Device": .29, "Quality": .14},
}
# Outcome allocation differs from population allocation; rate filters therefore
# divide the selected synthetic numerator by its selected synthetic denominator.
slice_skew = {
    "country": {"US": 1.01, "UK": .98, "DE": 1.04, "Other": .91},
    "appPlatform": {"iOS": 1.05, "Android": .96},
    "productLine": {"Bird": 1.03, "Hunting": .93},
    "deviceModel": {"K6": .89, "Bird Lite": .98, "Bird Pro": 1.08, "Hunt Pro": .96},
    "deviceStatus": {"Effective": 1.04, "Inactive": .64},
    "season": {"Migration": 1.04, "Other": .91},
    "salesChannel": {"Amazon": .98, "Shopify": 1.03},
    "subscriptionPlatform": {"App Store": 1.04, "Google Play": .93, "Web": 1.04},
    "plan": {"Free": .92, "Starter": 1.02, "Plus": 1.08, "Pro": 1.12},
    "billingCycle": {"monthly": 1.0, "annual": 1.0},
    "firmware": {"2.8": .82, "Other": 1.06},
    "functionType": {"Upload": 1.02, "Recognition": .94, "Push": 1.05},
    "anomalyType": {"Business": 1.09, "Device": 1.05, "Quality": .56},
}
slices = {}
for dimension, options in slice_weights.items():
    normalizer = sum(share * slice_skew[dimension][name] for name, share in options.items())
    slices[dimension] = {
        name: {"denominatorShare": share,
               "numeratorShare": round(share * slice_skew[dimension][name] / normalizer, 8)}
        for name, share in options.items()
    }
write("filter-slices.json", slices)
print(f"Generated {len(list(DATA.glob('*.json')))} JSON files and {len(values)} metrics")
